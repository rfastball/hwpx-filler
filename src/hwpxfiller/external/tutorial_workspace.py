"""튜토리얼 전용 저장소 — 학습 기록과 과정별 연습 홈(#1126).

연습은 사용자 환경에 아무것도 쓰지 않는다. 모든 연습 자원은 앱 홈 아래
``tutorial_workspace/`` 에만 생긴다:

- ``progress.json`` — 과정별 학습 위치·완료 기록(사용자 ``settings.json`` 이 아니다).
- ``lessons/<과정>/<회차>/`` — 그 과정 한 회차의 홈. 템플릿(``templates/``)·데이터(``data/``)·
  작업·데이터 풀·템플릿 권위·저작 기록·작업별 설정이 이 아래에 산다.

과정을 시작하거나 처음부터 할 때마다 **새 회차 홈**을 만든다. 이전 회차 홈은 지우려 시도하고,
지우지 못한 잔재도 이 저장소 안에만 남는다(사용자 환경으로 새지 않는다).
"""

from __future__ import annotations

import json
import shutil
import uuid
from pathlib import Path

from ..host.locations import tutorial_workspace_root
from . import settings
from .atomic import write_text_atomic
from .tutorial_practice import ORIGINALS, asset_root, fingerprint

__all__ = ["TutorialWorkspace", "DATA_NAME", "MODIFIED_REASON"]

#: 동봉 연습 데이터 — 원본 파생 사본도 이 파일에서 만든다.
DATA_NAME = "공고목록.xlsx"
_PROGRESS = "progress.json"
#: 설정 「튜토리얼 버튼 표시」(#1147) — 학습 기록과 다른 파일이라 「학습 기록 초기화」에 지워지지 않는다.
_ENTRY = "entry.json"
MODIFIED_REASON = "연습 파일이 수정됐습니다. 현재 파일을 보존합니다."


def _legacy_progress(raw: dict) -> dict:
    """#1126 이전 기록에서 위치·완료만 가져온다 — 옛 연습 사본 경로(문맥)는 쓰지 않는다."""
    records = raw.get("records") if isinstance(raw.get("records"), dict) else {}
    assert isinstance(records, dict)
    kept = {lesson: {"checkpoint": entry.get("checkpoint"), "completed": entry.get("completed")}
            for lesson, entry in records.items() if isinstance(entry, dict)}
    return {"version": raw.get("version"), "invite_seen": raw.get("invite_seen"),
            "selected": raw.get("selected"), "records": kept}


class TutorialWorkspace:
    """학습 기록 왕복과 과정별 연습 홈의 준비·확인."""

    def __init__(self, root: "Path | None" = None) -> None:
        self.root = Path(root) if root is not None else tutorial_workspace_root()

    # ------------------------------------------------------------ 학습 기록
    def load_progress(self) -> dict:
        path = self.root / _PROGRESS
        try:
            raw = json.loads(path.read_text(encoding="utf-8"))
        except FileNotFoundError:
            return _legacy_progress(settings.load_legacy_tutorial_lessons())
        except (OSError, ValueError) as exc:
            settings.alert(f"튜토리얼 학습 기록 판독 실패 — 처음 상태로 엽니다: {exc!r}")
            return {}
        return raw if isinstance(raw, dict) else {}

    def save_progress(self, value: dict) -> None:
        if not isinstance(value, dict) or value.get("version") != 1:
            raise ValueError("학습 기록 형식이 올바르지 않습니다.")
        self.root.mkdir(parents=True, exist_ok=True)
        write_text_atomic(self.root / _PROGRESS, json.dumps(value, ensure_ascii=False, indent=2))

    # ------------------------------------------------------------ 튜토리얼 버튼 표시(#1147)
    def load_entry_visible(self) -> bool:
        """미저장은 표시(기본)다. 손상은 알리고 표시로 연다 — 버튼을 잃으면 다시 켤 길이 설정뿐이다."""
        try:
            raw = json.loads((self.root / _ENTRY).read_text(encoding="utf-8"))
        except FileNotFoundError:
            return True
        except (OSError, ValueError) as exc:
            settings.alert(f"튜토리얼 버튼 표시 설정 판독 실패 — 표시로 엽니다: {exc!r}")
            return True
        visible = raw.get("visible") if isinstance(raw, dict) else None
        return visible if isinstance(visible, bool) else True

    def save_entry_visible(self, visible: bool) -> None:
        if not isinstance(visible, bool):
            raise ValueError("튜토리얼 버튼 표시 값이 올바르지 않습니다.")
        self.root.mkdir(parents=True, exist_ok=True)
        write_text_atomic(self.root / _ENTRY, json.dumps({"visible": visible}))

    # ------------------------------------------------------------ 연습 홈
    def seed(self, lesson_id: str, *, derived: str = "", derived_name: str = "") -> dict:
        """그 과정의 새 회차 홈을 만들고 동봉 원본을 원래 이름으로 놓는다.

        ``derived`` 는 원본과 다른 연습 데이터 한 벌(``replacement``·``blank``)을 더한다 —
        파일 이름은 과정 안내가 그대로 부를 수 있게 ``derived_name`` 으로 고정한다(#1127).
        """
        if derived not in {"", "blank", "replacement"} or bool(derived) != bool(derived_name):
            raise ValueError("알 수 없는 연습 데이터 종류입니다.")
        source = asset_root()
        missing = [name for name in ORIGINALS if not (source / name).is_file()]
        if missing:
            raise FileNotFoundError("동봉 예제를 찾을 수 없습니다: " + ", ".join(missing))
        batch = uuid.uuid4().hex[:10]
        home = self.root / "lessons" / lesson_id / batch
        try:
            assets = {name: self._place(source / name, home, name, batch) for name in ORIGINALS}
            extra = self._derive(source / DATA_NAME, home, batch, derived, derived_name) if derived else None
        except BaseException:
            shutil.rmtree(home, ignore_errors=True)
            raise
        return {"home": str(home), "batch": batch, "assets": assets, "derived_entry": extra}

    @staticmethod
    def _folder(home: Path, name: str) -> Path:
        return home / ("data" if name.endswith(".xlsx") else "templates")

    def _place(self, original: Path, home: Path, target_name: str, batch: str) -> dict:
        folder = self._folder(home, original.name)
        folder.mkdir(parents=True, exist_ok=True)
        target = folder / target_name
        with original.open("rb") as reader, target.open("xb") as writer:
            shutil.copyfileobj(reader, writer)
        return {"name": original.name, "path": str(target), "sha256": fingerprint(target),
                "batch": batch}

    def _derive(self, original: Path, home: Path, batch: str, derived: str, name: str) -> dict:
        if Path(name).name != name or Path(name).suffix != original.suffix or name == original.name:
            raise ValueError("연습 데이터 이름을 확인할 수 없습니다.")
        entry = self._place(original, home, name, batch)
        if derived == "blank":
            from openpyxl import load_workbook

            workbook = load_workbook(entry["path"])
            sheet = workbook["계약"]
            column = next(cell.column for cell in sheet[1] if cell.value == "계약보증금")
            # 계약보증금을 모든 행에서 비운다 — 작업대가 어느 행을 먼저 보이든 그 행에 〈빈 값〉이 선다.
            for row in range(2, sheet.max_row + 1):
                sheet.cell(row=row, column=column).value = None
            workbook.save(entry["path"])
            entry["sha256"] = fingerprint(Path(entry["path"]))
        entry["derived"] = derived
        return entry

    def validate(self, entry: dict, home: "str | Path") -> "tuple[bool, str]":
        """연습 홈 안의 기록된 파일이 그대로인가 — 다른 위치·수정·부재는 사유와 함께 거절."""
        try:
            name, path = entry["name"], Path(entry["path"])
            folder = self._folder(Path(home), name)
            if name not in ORIGINALS or path.is_symlink():
                return False, "연습 파일 경로를 확인할 수 없습니다."
            if path.resolve().parent != folder.resolve():
                return False, "연습 파일이 다른 위치를 가리킵니다."
            if not path.is_file():
                return False, "연습 파일이 없거나 이동했습니다."
            if fingerprint(path) != entry["sha256"]:
                return False, MODIFIED_REASON
            return True, ""
        except (KeyError, OSError, TypeError, ValueError):
            return False, "연습 파일 상태를 확인할 수 없습니다."

    def sweep(self, lesson_id: str, keep: "str | Path") -> None:
        """그 과정의 지난 회차 홈을 지운다 — 실패한 잔재는 이 저장소 안에만 남는다."""
        lessons = self.root / "lessons" / lesson_id
        keep_path = Path(keep).resolve()
        try:
            stale = [child for child in lessons.iterdir()
                     if child.is_dir() and not child.is_symlink() and child.resolve() != keep_path]
        except OSError:
            return
        for child in stale:
            shutil.rmtree(child, ignore_errors=True)
