"""「데이터로 필드 찾기」 인식률 — 누름틀 템플릿을 채우고 누름틀을 걷은 문서에서 원래 누름틀 자리를 되찾는가.

폐루프: 누름틀 템플릿을 데이터 행 하나로 채운다(``FieldDocument.set_field``) → 누름틀 제어(``fieldBegin``·
``fieldEnd`` 만 든 ``hp:ctrl``)를 걷어 평문 문서로 만든다 → ``read_hwpx`` → ``propose`` → 채우기 직후 문서의
누름틀 구간(정답)과 제안 자리를 대조한다. 금액 뒤 「원」까지 누름틀이고 제안 자리는 수만인 경우(``53,350,000``
⟷ ``53,350,000원``)는 같은 자리로 센다 — 제안은 단위를 값 표시에만 붙인다.

튜토리얼(「물품 구매입찰 공고」 + 공고목록 1행)은 39자리 전부, 시나리오·실제 공고는 이번 규칙이 닿은 수를
그대로 고정한다(빈 값 자리와, 일자·시각을 나눈 누름틀에 합친 일시 열이 겹치는 자리는 범위 밖).
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import openpyxl
from lxml import etree

from hwpxcore.package import HwpxPackage
from hwpxcore.text_extract import HP_NS
from hwpxfiller.domain.field_induction.proposal import propose
from hwpxfiller.domain.fields import FieldDocument, field_xml_names
from hwpxfiller.domain.mapping import FieldMapping, MappingProfile
from hwpxfiller.domain.schema import extract_schema, infer_type
from hwpxfiller.external.hwpx_field_proposal import read_hwpx

ROOT = Path(__file__).resolve().parents[1]
TUTORIAL = ROOT / "examples" / "tutorial"
SCENARIO = ROOT / "tests" / "corpus" / "scenario"
_HP = f"{{{HP_NS}}}"
_FIELD_CONTROL = {f"{_HP}fieldBegin", f"{_HP}fieldEnd"}
# 튜토리얼 seed 와 같은 열 이름 연결(템플릿 필드 → 데이터 열).
_SOURCE = {"낙찰자결정방법": "낙찰방법", "담당자 전화번호": "담당자전화", "대표계약업체": "계약상대자"}


@dataclass(frozen=True)
class Recall:
    gold: list[str]  # 누름틀 구간의 글자(문서 차례)
    missed: list[str]  # 제안으로 되찾지 못한 누름틀 구간의 글자
    false: list[str]  # 누름틀이 아니었던 제안 자리의 글자


def _mapping(fields: list[str]) -> MappingProfile:
    rows = []
    for name in fields:
        kind = infer_type(name) if infer_type(name) in {"date", "amount"} else "text"
        fmt = "ym" if name == "게시일시" and kind == "date" else ""
        rows.append(FieldMapping(name, _SOURCE.get(name, name), type=kind, fmt=fmt))
    return MappingProfile(mappings=rows)


def _rows(path: Path, sheet: str) -> tuple[list[str], list[dict[str, str]]]:
    book = openpyxl.load_workbook(path, read_only=True)
    lines = list(book[sheet].iter_rows(values_only=True))
    columns = [str(cell) for cell in lines[0]]
    return columns, [{column: "" if cell is None else str(cell) for column, cell in zip(columns, line, strict=False)}
                     for line in lines[1:]]


def _fill(template: bytes, values: dict[str, str]) -> HwpxPackage:
    package = HwpxPackage.from_bytes(template)
    for name in field_xml_names(package):
        document = FieldDocument(package.entries[name], entry=name)
        changed = False
        for field, value in values.items():
            changed |= document.set_field(field, value)
        if changed:
            package.entries[name] = document.to_bytes()
    return package


def _strip_fields(package: HwpxPackage) -> HwpxPackage:
    """누름틀 제어만 걷는다 — 채운 글자는 그대로 남아 평문 문서가 된다."""
    out = HwpxPackage.from_bytes(package.to_bytes())
    parser = etree.XMLParser(remove_blank_text=False, resolve_entities=False)
    for entry in out.content_xml_names():
        root = etree.fromstring(out.entries[entry], parser=parser)
        for control in list(root.iter(f"{_HP}ctrl")):
            if len(control) and all(child.tag in _FIELD_CONTROL for child in control):
                _drop_keeping_tail(control)
        out.entries[entry] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    return out


def _drop_keeping_tail(node: etree._Element) -> None:
    parent, previous = node.getparent(), node.getprevious()
    assert parent is not None
    if node.tail:
        if previous is not None:
            previous.tail = (previous.tail or "") + node.tail
        else:
            parent.text = (parent.text or "") + node.tail
    parent.remove(node)


def recall(template: Path, data: Path, sheet: str, row: int) -> Recall:
    """``row``(0부터) 행으로 채운 ``template`` 의 누름틀 자리를 제안이 얼마나 되찾는가."""
    raw = template.read_bytes()
    columns, records = _rows(data, sheet)
    filled = _fill(raw, _mapping([field.name for field in extract_schema(HwpxPackage.from_bytes(raw)).fields])
                   .apply(records[row]))
    truth, _ = read_hwpx(filled)
    paragraphs, _ = read_hwpx(_strip_fields(filled))
    assert [item.text for item in paragraphs] == [item.text for item in truth]
    gold = {(index, lo, hi): item.text[lo:hi] for index, item in enumerate(truth) for lo, hi in item.field_spans}
    found: set[tuple[int, int, int]] = set()
    false: list[str] = []
    for group in propose(paragraphs, records[row], columns).groups:
        for spot in group.spots if group.kind == "proposal" else ():
            key = next((place for place in ((spot.paragraph, spot.start, spot.end),
                                            (spot.paragraph, spot.start, spot.end + 1))
                        if gold.get(place) in (spot.text, spot.text + "원")), None)
            if key is None:
                false.append(spot.text)
            else:
                found.add(key)
    return Recall(list(gold.values()), [text for key, text in gold.items() if key not in found], false)


def test_tutorial_notice_row_one_recovers_every_field_place() -> None:
    result = recall(TUTORIAL / "물품 구매입찰 공고.hwpx", TUTORIAL / "공고목록.xlsx", "공고", 0)
    assert len(result.gold) == 39
    assert result.missed == []
    # 조달청 담당자 줄(「주무관 홍길동」)의 이름은 템플릿에 박힌 글자다 — 제안 자리가 아니다.
    assert result.false == []


def test_scenario_bid_notice_row_one() -> None:
    result = recall(SCENARIO / "templates" / "입찰공고서.hwpx", SCENARIO / "data" / "조달_한글.xlsx", "Sheet", 0)
    assert len(result.gold) == 25
    # 남는 것: 빈 값 누름틀 2곳, 입찰마감 일자·시각을 나눈 누름틀(데이터는 합친 일시 열 — 범위 밖).
    assert sorted(result.missed) == ["", "", "17:00", "2026. 7. 21."]
    assert result.false == []


def test_real_limited_notice_row_one() -> None:
    result = recall(ROOT / "tests" / "corpus" / "real" / "bid_notice_limited_under100m.hwpx",
                    SCENARIO / "data" / "조달_한글.xlsx", "Sheet", 0)
    assert len(result.gold) == 27
    assert result.missed == [""] * 5  # 빈 값 누름틀뿐
    # 「입찰방법:」 라벨이 있는 문서의 「입찰방식: 전자입찰(국내입찰)」은 다른 항목이다.
    assert result.false == []
